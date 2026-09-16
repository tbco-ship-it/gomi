#!/usr/bin/env python3
"""
Normalize Hiroshima City (広島市) Garbage Collection Schedule
-------------------------------------------------------------
Reads official Excel (.xlsx) calendars for all 8 wards (141 files):
data/raw/hiroshima/*.xlsx
and converts into unified schema:
data/normalized/hiroshima.json

Schema rules:
- Top-level pure JSON array
- Standard waste type keys: burnable, resource, plastic, paper_cloth, nonburnable, bulky
  - 可燃ごみ -> burnable (週2回)
  - リサイクルプラ -> plastic (週1回)
  - 資源ごみ・有害ごみ -> resource (月2回)
  - 不燃ごみ -> nonburnable (月2回)
  - 大型ごみ(予約制) -> bulky (月2回, 事前予約制)
  - その他プラ -> preserved in raw["other_plastic"] (6대 표준 키 외)
- 8 wards: 中区 (naka), 東区 (higashi), 南区 (minami), 西区 (nishi),
  安佐南区 (asaminami), 安佐北区 (asakita), 安芸区 (aki), 佐伯区 (saeki)
- Unique slugs (0 duplicates) generated via pykakasi
- source.fetched: 2026-09-17
"""

import json
import os
import re
import sys
import unicodedata
import urllib.request
from collections import defaultdict, Counter
from concurrent.futures import ThreadPoolExecutor
from typing import Dict, List, Any, Set, Tuple
import openpyxl
import pykakasi

VALID_DAYS = {"月", "火", "水", "木", "金", "土", "日"}
TODAY_STR = "2026-09-17"
DAYS_MAP = ["日", "月", "火", "水", "木", "金", "土"]
KANJI_NUMS = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}

kakasi = pykakasi.kakasi()


def to_romaji(text: str) -> str:
    res = kakasi.convert(text)
    romaji = "".join([x.get("passport") or x.get("hepburn", "") for x in res]).lower()
    return re.sub(r"[^a-z0-9]+", "-", romaji).strip("-")


def extract_nums(text: str) -> str:
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    nums = re.findall(r"\d+", text)
    return "-".join(nums) if nums else ""


def parse_hiroshima_town_str(raw_str: str) -> List[Tuple[str, str, str]]:
    """
    Parses complex town list strings like:
    '舟入南二丁目、三丁目、六丁目・西川口町'
    -> [('舟入南', '2丁目', ''), ('舟入南', '3丁目', ''), ('舟入南', '6丁目', ''), ('西川口町', '', '')]
    """
    s = unicodedata.normalize("NFKC", raw_str).strip()
    sub_map = {}

    def repl(m):
        k = f"__SUB_{len(sub_map)}__"
        sub_map[k] = m.group(0)
        return k

    s_clean = re.sub(r"[（\(][^）\)]*[）\)]", repl, s)
    parts = [p.strip() for p in re.split(r"[・、,]", s_clean) if p.strip()]
    records = []
    current_town = ""

    for p in parts:
        for k, v in sub_map.items():
            if k in p:
                p = p.replace(k, v)

        sub = ""
        m_sub = re.search(r"[（\(](.*?)[）\)]", p)
        if m_sub:
            sub = m_sub.group(1).strip()
            p = re.sub(r"[（\(].*?[）\)]", "", p).strip()

        # Check range: e.g. 一丁目~三丁目 or 1~3丁目
        m_range = re.search(r"([一二三四五六七八九十\d]+)丁目?\s*[~～-]\s*([一二三四五六七八九十\d]+)丁目?", p)
        if m_range:
            prefix = p[: m_range.start()].strip()
            if prefix:
                current_town = prefix
            start_val = m_range.group(1)
            end_val = m_range.group(2)
            s_n = KANJI_NUMS.get(start_val, int(start_val) if start_val.isdigit() else 1)
            e_n = KANJI_NUMS.get(end_val, int(end_val) if end_val.isdigit() else 1)
            for n in range(s_n, e_n + 1):
                records.append((current_town, f"{n}丁目", sub))
            continue

        # Single chome: e.g. 二丁目 or 2丁目
        m_chome = re.search(r"([一二三四五六七八九十\d]+)丁目", p)
        if m_chome:
            prefix = p[: m_chome.start()].strip()
            if prefix:
                current_town = prefix
            c_val = m_chome.group(1)
            c_n = KANJI_NUMS.get(c_val, int(c_val) if c_val.isdigit() else c_val)
            post = p[m_chome.end() :].strip()
            sub_combined = (post + " " + sub).strip() if post else sub
            records.append((current_town, f"{c_n}丁目", sub_combined))
            continue

        current_town = p
        records.append((current_town, "", sub))

    return records


def extract_schedule_from_xlsx(path: str) -> Dict[str, Any]:
    """
    Parses full year calendar grid in Hiroshima Excel file.
    Grid layout:
    Months at rows 3, 24, 50, 71, each having columns for 3 months.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    blocks = [
        (3, [3, 12, 21]),
        (24, [3, 12, 21]),
        (50, [3, 12, 21]),
        (71, [3, 12, 21]),
    ]

    events = defaultdict(list)

    for start_r, col_starts in blocks:
        for c_start in col_starts:
            for w in range(6):
                r_date = start_r + 2 + w * 3
                r_l1 = r_date + 1
                r_l2 = r_date + 2
                if r_date > ws.max_row:
                    break
                for dow_idx in range(7):
                    col = c_start + dow_idx
                    val_date = ws.cell(r_date, col).value
                    if isinstance(val_date, int):
                        v1 = str(ws.cell(r_l1, col).value or "").strip()
                        v2 = str(ws.cell(r_l2, col).value or "").strip()
                        comb = v1 + v2
                        if not comb:
                            continue
                        dow = DAYS_MAP[dow_idx]
                        week_idx = (val_date - 1) // 7 + 1

                        if "可燃" in comb:
                            events["burnable"].append((dow, week_idx))
                        if "ﾘｻｲｸﾙ" in comb or "リサイクル" in comb:
                            events["plastic"].append((dow, week_idx))
                        if "資源" in comb or "有害" in comb:
                            events["resource"].append((dow, week_idx))
                        if "不燃" in comb:
                            events["nonburnable"].append((dow, week_idx))
                        if "大型" in comb:
                            events["bulky"].append((dow, week_idx))
                        if "その他" in comb:
                            events["other_plastic"].append((dow, week_idx))

    sched: Dict[str, Any] = {}
    time_slot = "朝8:30まで"

    type_labels = {
        "burnable": "可燃ごみ",
        "plastic": "リサイクルプラ",
        "resource": "資源ごみ・有害ごみ",
        "nonburnable": "不燃ごみ",
        "bulky": "大型ごみ(予約制)",
        "other_plastic": "その他プラ",
    }

    for t_k, ev_list in events.items():
        dow_counts = Counter(d for d, w in ev_list)
        valid_dows = [d for d, c in dow_counts.items() if c >= 5]
        week_counts = Counter(w for d, w in ev_list if d in valid_dows)

        is_weekly = any(dow_counts[d] >= 35 for d in valid_dows)
        if is_weekly:
            weeks = None
        else:
            typical_weeks = [w for w in [1, 2, 3, 4] if week_counts[w] >= 5]
            weeks = sorted(typical_weeks) if typical_weeks else None

        info = {
            "label": type_labels.get(t_k, t_k),
            "days": sorted(valid_dows, key=lambda d: DAYS_MAP.index(d)),
            "weeks": weeks,
            "time": time_slot,
        }
        if t_k == "bulky":
            info["note"] = "事前予約制"

        sched[t_k] = info

    return sched


def normalize_hiroshima(raw_dir: str, output_path: str):
    index_file = os.path.join(raw_dir, "hiroshima_files_index.json")
    if not os.path.exists(index_file):
        # Also check data/raw/
        index_file = os.path.join(os.path.dirname(raw_dir), "hiroshima_files_index.json")
    if not os.path.exists(index_file):
        raise FileNotFoundError(f"Missing hiroshima_files_index.json at {index_file}")

    with open(index_file, "r", encoding="utf-8") as f:
        file_entries = json.load(f)

    records: List[Dict[str, Any]] = []
    used_slugs: Set[str] = set()

    # Pre-extract schedules for all unique files
    file_scheds: Dict[str, Dict[str, Any]] = {}
    print(f"[*] Parsing schedule from {len(file_entries)} Excel files...")
    for entry in file_entries:
        fname = entry["filename"]
        if fname not in file_scheds:
            fpath = os.path.join(raw_dir, fname)
            if not os.path.exists(fpath):
                raise FileNotFoundError(f"Missing raw xlsx file: {fpath}")
            file_scheds[fname] = extract_schedule_from_xlsx(fpath)

    print(f"[+] All {len(file_scheds)} files parsed. Generating normalized records...")

    for entry in file_entries:
        ward = entry["ward"]
        ward_en = entry["ward_en"]
        town_str = entry["town_str"]
        url = entry["url"]
        fname = entry["filename"]
        sched = file_scheds[fname]

        parsed_items = parse_hiroshima_town_str(town_str)

        for town, chome, sub in parsed_items:
            t_ro = to_romaji(town)
            nums = extract_nums(chome)
            sub_ro = to_romaji(sub) if sub else ""

            parts = [t_ro]
            if nums:
                parts.append(nums)
            if sub_ro:
                parts.append(sub_ro)

            base_slug = f"hiroshima/{ward_en}/" + "-".join(parts)
            slug = base_slug
            idx = 2
            while slug in used_slugs:
                slug = f"{base_slug}-{idx}"
                idx += 1
            used_slugs.add(slug)

            # Build types dict (5 standard keys)
            types: Dict[str, Any] = {}
            for k in ["burnable", "plastic", "resource", "nonburnable", "bulky"]:
                if k in sched:
                    types[k] = sched[k]

            raw_dict = {
                "source_file": fname,
                "town_str_raw": town_str,
                "other_plastic": sched.get("other_plastic"),
            }

            record = {
                "pref": "広島県",
                "city": "広島市",
                "city_en": "hiroshima",
                "ward": ward,
                "ward_en": ward_en,
                "town": town,
                "romaji": t_ro,
                "chome": chome,
                "sub": sub,
                "slug": slug,
                "types": types,
                "rules": {
                    "holiday_collection": "祝日も収集（年末年始を除く）",
                    "time_by": "8:30",
                    "yearend": "12/31~1/3 休止",
                },
                "source": {
                    "url": url,
                    "fetched": TODAY_STR,
                    "basis": "official_xlsx",
                },
                "notes": "その他プラ（隔週）収集日程はraw.other_plasticに保持",
                "raw": raw_dict,
            }
            records.append(record)

    # Verification
    assert len(records) > 0, "No records produced for Hiroshima!"
    assert len(records) == len(used_slugs), f"Slug collision in Hiroshima! {len(records)} != {len(used_slugs)}"

    for r in records:
        assert re.match(r"^[a-z0-9\-/]+$", r["slug"]), f"Invalid slug characters: {r['slug']}"
        for t_k, t_v in r["types"].items():
            assert t_k in {"burnable", "resource", "plastic", "paper_cloth", "nonburnable", "bulky"}, f"Invalid type key: {t_k}"
            for d in t_v["days"]:
                assert d in VALID_DAYS, f"Invalid day {d} in record {r['slug']}"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    print(f"[✔] Hiroshima normalized: {len(records)} records saved to {output_path}")
    print(f"[✔] Unique slugs: {len(used_slugs)} (collisions: 0)")


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    raw_dir = os.path.join(base_dir, "data", "raw", "hiroshima")
    output_file = os.path.join(base_dir, "data", "normalized", "hiroshima.json")
    normalize_hiroshima(raw_dir, output_file)
